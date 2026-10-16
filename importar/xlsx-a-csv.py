# ============================================================
# VetaApp · importar/xlsx-a-csv.py
# Convierte el .xlsx que baja Google Sheets en el CSV que lee armar-sql.mjs.
#
# ---------- POR QUÉ EXISTE ESTE PASO ----------
# Se podría bajar el CSV directo de Google y saltarse esto. No conviene:
# el .xlsx guarda los TIPOS de cada celda, y el CSV los pierde. Con los
# tipos a la vista, tres datos que en el CSV serían una adivinanza quedan
# resueltos de entrada.
#
#   · LAS FECHAS son `datetime`, no texto. En el CSV de Google salen como
#     "26/07/2018" o "2018-07-26" según el idioma de la cuenta, y hay que
#     adivinar cuál número es el día. Acá no hay nada que adivinar.
#
#   · LA CÉDULA es `float`: 1000000200.0. Escrita en crudo queda
#     "1000000200.0", y el importador —que normaliza a solo dígitos— la
#     lee como 10000002000. Un cero de más. Esa cédula no identifica a
#     nadie, y el índice único de 011 la deja pasar porque es otro texto.
#
#   · EL COSTO también es `float`: 40000.0 → "400000". Diez veces más
#     caro, en la columna de la que vive la peluquería.
#
# Los tres errores son silenciosos: nada falla, los datos quedan mal.
#
# ---------- CÓMO SE USA ----------
#   python3 importar/xlsx-a-csv.py "~/Downloads/Base Clientes.xlsx" "Respuestas de formulario 1"
#
# Sin el nombre de la hoja, lista las que haya y no escribe nada.
# ============================================================

import csv
import sys
from datetime import datetime, date, time

try:
    import openpyxl
except ImportError:
    sys.exit('Falta openpyxl: pip3 install --user openpyxl')


def texto(celda):
    """Una celda de Excel al texto que va en el CSV, sin inventar ni perder."""
    if celda is None:
        return ''

    # Fecha con hora (la marca temporal del formulario): se escribe
    # completa. `armar-sql.mjs` parte por el espacio para la fecha y busca
    # HH:MM aparte para `tratamientos.hora`, que es la hora real a la que
    # la clínica atendió.
    if isinstance(celda, datetime):
        if (celda.hour, celda.minute, celda.second) == (0, 0, 0):
            return celda.strftime('%Y-%m-%d')
        return celda.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(celda, date):
        return celda.strftime('%Y-%m-%d')
    if isinstance(celda, time):
        return celda.strftime('%H:%M:%S')

    # El entero disfrazado de float: el caso de la cédula y del costo.
    # `is_integer()` y no `int(x)` a secas porque un precio con centavos
    # —no hay ninguno hoy, pero puede llegar— no se debe truncar en
    # silencio: se escribe tal cual y el importador decide.
    if isinstance(celda, float):
        return str(int(celda)) if celda.is_integer() else repr(celda)

    if isinstance(celda, bool):
        return 'Si' if celda else 'No'

    return str(celda).strip()


def main():
    if len(sys.argv) < 2:
        sys.exit('Uso: python3 importar/xlsx-a-csv.py <archivo.xlsx> [nombre de la hoja]')

    ruta = sys.argv[1]
    # data_only=True entrega el VALOR de las fórmulas, no la fórmula. Si la
    # hoja tuviera un =SUMA(), sin esto llegaría el texto "=SUMA(...)".
    libro = openpyxl.load_workbook(ruta, data_only=True)

    if len(sys.argv) < 3:
        print('Hojas en el archivo:')
        for h in libro.worksheets:
            print(f'  · {h.title}')
        sys.exit('\nElegí una y volvé a correr esto con el nombre como segundo argumento.')

    nombre = sys.argv[2]
    if nombre not in libro.sheetnames:
        sys.exit(f'No existe la hoja {nombre!r}. Hay: {", ".join(libro.sheetnames)}')

    hoja = libro[nombre]
    filas = []
    for cruda in hoja.iter_rows(values_only=True):
        fila = [texto(c) for c in cruda]
        # Las hojas de Google vienen estiradas hasta la fila 155 con filas
        # en blanco. Se cortan acá y no en el importador para que el conteo
        # que imprime al final sea el de verdad.
        if any(fila):
            filas.append(fila)

    salida = sys.argv[3] if len(sys.argv) > 3 else 'importar/base.csv'
    # newline='' es de la documentación de csv: sin eso, en Windows cada
    # fila sale con un salto de línea de más.
    with open(salida, 'w', newline='', encoding='utf-8') as f:
        csv.writer(f).writerows(filas)

    print(f'{len(filas) - 1} filas con datos (sin contar el encabezado) → {salida}')


if __name__ == '__main__':
    main()
